#!/usr/bin/env python3
"""Extrait les couleurs (fond, texte, gras) du plan de test Excel vers un JSON
compagnon consommé par plan-de-test.js.

SheetJS (édition communautaire) ne lit que les fonds, pas la couleur de texte
ni le gras : on pré-calcule donc ces styles ici avec openpyxl. On ne stocke que
les cellules qui diffèrent du style par défaut (le plus fréquent) -> JSON léger.

À relancer chaque fois que le .xlsx change :
    python frontend/scripts/gen-plan-styles.py
"""
import json
import os
from collections import Counter

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PUBLIC = os.path.normpath(os.path.join(HERE, "..", "public"))
XLSX = os.path.join(PUBLIC, "Plan_de_test_MoodIT.xlsx")
OUT = os.path.join(PUBLIC, "plan-de-test.styles.json")


def hex6(color):
    """Renvoie 'RRGGBB' pour une couleur RGB openpyxl, ou None (couleur de
    thème/indexée non résolue -> on laissera le défaut CSS s'appliquer)."""
    rgb = getattr(color, "rgb", None)
    if not isinstance(rgb, str):
        return None
    rgb = rgb[2:] if len(rgb) == 8 else rgb  # retire l'alpha 'FF'
    return rgb.upper() if len(rgb) == 6 else None


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # 1re passe : déterminer le fond et la couleur de texte par défaut (majoritaires).
    bg_count, fg_count = Counter(), Counter()
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if c.value is None and not (c.fill and c.fill.patternType):
                    continue
                if c.fill and c.fill.patternType:
                    b = hex6(c.fill.fgColor)
                    if b:
                        bg_count[b] += 1
                if c.font:
                    f = hex6(c.font.color)
                    if f:
                        fg_count[f] += 1

    def_bg = bg_count.most_common(1)[0][0] if bg_count else "FFFFFF"
    def_fg = fg_count.most_common(1)[0][0] if fg_count else "000000"

    # 2e passe : ne garder que les cellules qui s'écartent du défaut.
    sheets = {}
    for name in wb.sheetnames:
        cells = {}
        for row in wb[name].iter_rows():
            for c in row:
                if c.value is None and not (c.fill and c.fill.patternType):
                    continue
                b = hex6(c.fill.fgColor) if (c.fill and c.fill.patternType) else None
                f = hex6(c.font.color) if c.font else None
                bold = bool(c.font and c.font.bold)
                entry = {}
                if b and b != def_bg:
                    entry["b"] = b
                if f and f != def_fg:
                    entry["f"] = f
                if bold:
                    entry["w"] = 1
                if entry:
                    # clés 0-based (r,c) pour coller au décodage SheetJS côté client
                    cells["%d,%d" % (c.row - 1, c.column - 1)] = entry
        if cells:
            sheets[name] = cells

    out = {"defaults": {"bg": def_bg, "fg": def_fg}, "sheets": sheets}
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(",", ":"))

    overrides = sum(len(v) for v in sheets.values())
    size = os.path.getsize(OUT)
    print("defaut  fond #%s / texte #%s" % (def_bg, def_fg))
    print("cellules stylees (hors defaut) : %d" % overrides)
    print("ecrit %s (%.1f Ko)" % (OUT, size / 1024))


if __name__ == "__main__":
    main()
